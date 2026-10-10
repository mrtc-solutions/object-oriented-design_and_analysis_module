"""Build printable DOCX and PDF copies of the end-of-semester papers and guide.

Run from the repository root with:
    .venv/bin/python tools/build_exam_docs.py

The Markdown files in exams/ remain the editable source of truth. Generated
DOCX/PDF files are written to output/exams/ for printing and distribution.
"""
from __future__ import annotations

import html
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (HRFlowable, Paragraph, SimpleDocTemplate,
                                Spacer, Table, TableStyle)

ROOT = Path(__file__).resolve().parents[1]
EXAMS = ROOT / "exams"
EXAM_OUTPUT = ROOT / "output" / "exams"
MAROON = "7B1113"
DARK = "59090C"
GOLD = "C9A227"
GREY = "6B6B6B"
INK = "2B2626"
WHITE = "FFFFFF"

PAPER_SOURCES = [
    EXAMS / "End-Semester-Exam-Paper-A.md",
    EXAMS / "End-Semester-Exam-Paper-B.md",
    EXAMS / "Examiner-Marking-Guide.md",
]

INLINE_TOKEN = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*]+\*)")
TABLE_SEPARATOR = re.compile(r"^:?-{3,}:?$")


def inline_tokens(text: str):
    """Yield (text, kind) pairs for the small Markdown inline subset in these files."""
    for token in INLINE_TOKEN.split(text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            yield token[2:-2], "bold"
        elif token.startswith("`") and token.endswith("`"):
            yield token[1:-1], "code"
        elif token.startswith("*") and token.endswith("*"):
            yield token[1:-1], "italic"
        else:
            yield token, "plain"


def reportlab_inline(text: str) -> str:
    parts = []
    for value, kind in inline_tokens(text):
        escaped = html.escape(value, quote=False)
        if kind == "bold":
            parts.append(f"<b>{reportlab_inline(value)}</b>")
        elif kind == "italic":
            parts.append(f"<i>{reportlab_inline(value)}</i>")
        elif kind == "code":
            parts.append(f'<font name="ExamMono">{escaped}</font>')
        else:
            parts.append(escaped)
    return "".join(parts)


def add_docx_inline(paragraph, text: str, *, white_bold: bool = False) -> None:
    for value, kind in inline_tokens(text):
        if kind == "bold" and "`" in value:
            start = len(paragraph.runs)
            add_docx_inline(paragraph, value, white_bold=white_bold)
            for run in paragraph.runs[start:]:
                run.bold = True
            continue
        run = paragraph.add_run(value)
        if kind == "bold" or white_bold:
            run.bold = True
        if kind == "italic":
            run.italic = True
        if kind == "code":
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        if white_bold:
            run.font.color.rgb = RGBColor.from_string(WHITE)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top=85, start=100, bottom=85, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "5")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "D8C9C9")


def md_table_rows(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if cells and all(TABLE_SEPARATOR.match(cell.replace(" ", "")) for cell in cells):
            continue
        rows.append(cells)
    return rows


def docx_table(document, rows: list[list[str]]) -> None:
    if not rows:
        return
    columns = max(len(row) for row in rows)
    table = document.add_table(rows=0, cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    section = document.sections[-1]
    available = section.page_width - section.left_margin - section.right_margin
    ratios = {2: (0.31, 0.69), 3: (0.22, 0.25, 0.53)}
    col_widths = tuple(int(available * ratio) for ratio in
                       ratios.get(columns, (1 / columns,) * columns))
    for column, width in zip(table.columns, col_widths):
        column.width = width
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for col in range(columns):
            cell = cells[col]
            cell.width = col_widths[col]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.space_before = Pt(0)
            value = values[col] if col < len(values) else ""
            if row_index == 0:
                shade_cell(cell, MAROON)
                add_docx_inline(paragraph, value, white_bold=True)
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)
            else:
                if row_index % 2 == 0:
                    shade_cell(cell, "F8F3F0")
                add_docx_inline(paragraph, value)
                for run in paragraph.runs:
                    run.font.size = Pt(8.7)
    set_table_borders(table)
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def add_horizontal_rule(document) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(8)
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "10")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), GOLD)
    borders.append(bottom)
    p_pr.append(borders)


def page_field(paragraph, code: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = code
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])


def is_candidate(path: Path) -> bool:
    return "Paper-" in path.name


def header_label(path: Path) -> str:
    if path.name.endswith("Paper-A.md"):
        return "PAPER A · CANDIDATE COPY"
    if path.name.endswith("Paper-B.md"):
        return "PAPER B · CANDIDATE COPY"
    return "CONFIDENTIAL · EXAMINER COPY"


def docx_build(source: Path, destination: Path) -> None:
    guide = "Guide" in source.name
    document = Document()
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    for style_name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11)):
        style = document.styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(MAROON)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12 if style_name != "Heading 1" else 16)
        style.paragraph_format.space_after = Pt(5)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run(f"EXPLOITS UNIVERSITY  |  BICT 3202  |  {header_label(source)}")
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(MAROON)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("BICT 3202 · Object-Oriented Analysis and Design  |  Page ").font.size = Pt(8)
    page_field(footer, "PAGE")
    footer.add_run(" of ").font.size = Pt(8)
    page_field(footer, "NUMPAGES")
    for run in footer.runs:
        run.font.size = Pt(8)
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    document.settings.element.append(update_fields)

    document.core_properties.title = source.stem.replace("-", " ")
    document.core_properties.subject = "BICT 3202 end-of-semester examination material"
    document.core_properties.author = "Exploits University — BICT 3202"
    if guide:
        document.core_properties.comments = "CONFIDENTIAL — examiner/lecturer use only; not for candidate distribution."

    lines = source.read_text(encoding="utf-8").splitlines()
    index = 0
    code_block = False
    while index < len(lines):
        raw = lines[index]
        line = raw.strip()
        if line.startswith("```"):
            code_block = not code_block
            index += 1
            continue
        if code_block:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.25)
            paragraph.paragraph_format.space_after = Pt(2)
            run = paragraph.add_run(raw)
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
            index += 1
            continue
        if not line:
            index += 1
            continue
        if line == "---":
            add_horizontal_rule(document)
            index += 1
            continue
        if line.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            docx_table(document, md_table_rows(table_lines))
            continue
        heading_match = re.match(r"^(#{1,4})\s+(.*)$", line)
        if heading_match:
            level = min(len(heading_match.group(1)), 3)
            heading = heading_match.group(2).strip()
            paragraph = document.add_paragraph(style=f"Heading {level}")
            add_docx_inline(paragraph, heading)
            if heading == "EXPLOITS UNIVERSITY" or heading.startswith("BICT 3202 — Examiner"):
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.size = Pt(18)
            index += 1
            continue
        if line.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.space_after = Pt(3)
            add_docx_inline(paragraph, line[2:])
            index += 1
            continue
        if re.match(r"^\d+\.\s+", line):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.22)
            paragraph.paragraph_format.first_line_indent = Inches(-0.22)
            paragraph.paragraph_format.space_after = Pt(4)
            add_docx_inline(paragraph, line)
            index += 1
            continue
        if line.startswith("> "):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.25)
            paragraph.paragraph_format.right_indent = Inches(0.25)
            add_docx_inline(paragraph, line[2:])
            index += 1
            continue
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.keep_together = True
        if guide and line.startswith("**("):
            paragraph.paragraph_format.keep_with_next = True
        if "END-OF-SEMESTER EXAMINATION" in line or line.startswith("**Confidential:"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(8)
        add_docx_inline(paragraph, line)
        if "END-OF-SEMESTER EXAMINATION" in line or "Confidential:" in line:
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor.from_string(MAROON if not guide else DARK)
                run.font.size = Pt(11)
        index += 1

    document.save(destination)


def setup_pdf_fonts() -> None:
    fonts = ROOT / "assets" / "fonts"
    font_specs = {
        "ExamBody": fonts / "OpenSans-Regular.ttf",
        "ExamBody-Bold": fonts / "OpenSans-Bold.ttf",
        "ExamBody-Italic": fonts / "OpenSans-Italic.ttf",
        "ExamHeading": fonts / "Poppins-Regular.ttf",
        "ExamHeading-Bold": fonts / "Poppins-Bold.ttf",
        "ExamMono": fonts / "DejaVuSansMono.ttf",
    }
    for name, path in font_specs.items():
        pdfmetrics.registerFont(TTFont(name, str(path)))
    pdfmetrics.registerFontFamily("ExamBody", normal="ExamBody", bold="ExamBody-Bold",
                                  italic="ExamBody-Italic", boldItalic="ExamBody-Bold")
    pdfmetrics.registerFontFamily("ExamHeading", normal="ExamHeading", bold="ExamHeading-Bold",
                                  italic="ExamHeading", boldItalic="ExamHeading-Bold")


def pdf_table(rows: list[list[str]], width: float):
    if not rows:
        return Spacer(1, 1)
    columns = max(len(row) for row in rows)
    if columns == 2:
        widths = [width * 0.31, width * 0.69]
    elif columns == 3:
        widths = [width * 0.22, width * 0.25, width * 0.53]
    else:
        widths = [width / columns] * columns
    cell_styles = [
        ParagraphStyle("table_header", fontName="ExamBody-Bold", fontSize=8.0,
                       leading=10, textColor=colors.white, alignment=TA_LEFT),
        ParagraphStyle("table_cell", fontName="ExamBody", fontSize=7.7,
                       leading=9.6, textColor=colors.HexColor("#" + INK), alignment=TA_LEFT),
    ]
    data = []
    for row_index, row in enumerate(rows):
        style = cell_styles[0] if row_index == 0 else cell_styles[1]
        data.append([Paragraph(reportlab_inline(row[col] if col < len(row) else ""), style)
                     for col in range(columns)])
    table = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
    commands = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + MAROON)),
        ("GRID", (0, 0), (-1, -1), 0.45, colors.HexColor("#D8C9C9")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]
    for row_index in range(2, len(rows), 2):
        commands.append(("BACKGROUND", (0, row_index), (-1, row_index), colors.HexColor("#F8F3F0")))
    table.setStyle(TableStyle(commands))
    return table


def pdf_build(source: Path, destination: Path) -> None:
    setup_pdf_fonts()
    guide = "Guide" in source.name
    paper = "A" if source.name.endswith("Paper-A.md") else "B" if source.name.endswith("Paper-B.md") else ""
    page_width, page_height = A4
    left = right = 48
    top = 54
    bottom = 44
    usable = page_width - left - right

    body = ParagraphStyle("exam_body", fontName="ExamBody", fontSize=8.9 if guide else 9.2,
                          leading=12.2 if guide else 13.2,
                          textColor=colors.HexColor("#" + INK), spaceAfter=4.5 if guide else 6,
                          alignment=TA_LEFT, allowWidows=0, allowOrphans=0)
    bullet = ParagraphStyle("exam_bullet", parent=body, leftIndent=15, firstLineIndent=-9,
                            spaceAfter=2.5 if guide else 4)
    ordered = ParagraphStyle("exam_ordered", parent=body, leftIndent=20, firstLineIndent=-20,
                             spaceAfter=3 if guide else 4)
    h1 = ParagraphStyle("exam_h1", fontName="ExamHeading-Bold", fontSize=15.5, leading=19,
                        textColor=colors.HexColor("#" + MAROON), alignment=TA_CENTER,
                        spaceBefore=7, spaceAfter=8, keepWithNext=True)
    h2 = ParagraphStyle("exam_h2", fontName="ExamHeading-Bold", fontSize=11.8 if guide else 12.2,
                        leading=14.4 if guide else 15,
                        textColor=colors.HexColor("#" + MAROON), spaceBefore=8 if guide else 10,
                        spaceAfter=4 if guide else 5, keepWithNext=True)
    h3 = ParagraphStyle("exam_h3", fontName="ExamHeading-Bold", fontSize=10.2 if guide else 10.5,
                        leading=12.2 if guide else 13,
                        textColor=colors.HexColor("#" + DARK), spaceBefore=6 if guide else 8,
                        spaceAfter=3 if guide else 4, keepWithNext=True)
    h4 = ParagraphStyle("exam_h4", fontName="ExamHeading-Bold", fontSize=9.4 if guide else 9.7,
                        leading=11.5 if guide else 12,
                        textColor=colors.HexColor("#" + MAROON), spaceBefore=5 if guide else 6,
                        spaceAfter=2.5 if guide else 3, keepWithNext=True)
    centered = ParagraphStyle("exam_centered", parent=body, alignment=TA_CENTER,
                              fontName="ExamBody-Bold", textColor=colors.HexColor("#" + MAROON),
                              spaceAfter=7)
    code_style = ParagraphStyle("exam_code", parent=body, fontName="ExamMono", fontSize=8.2,
                                leading=11, leftIndent=12, backColor=colors.HexColor("#F8F3F0"),
                                borderPadding=4)
    guide_lead = ParagraphStyle("guide_lead", parent=body, keepWithNext=True)
    story = []
    lines = source.read_text(encoding="utf-8").splitlines()
    index = 0
    code_block = False
    while index < len(lines):
        raw = lines[index]
        line = raw.strip()
        if line.startswith("```"):
            code_block = not code_block
            index += 1
            continue
        if code_block:
            story.append(Paragraph(html.escape(raw, quote=False) or "&nbsp;", code_style))
            index += 1
            continue
        if not line:
            index += 1
            continue
        if line == "---":
            story.append(Spacer(1, 3))
            story.append(HRFlowable(width="100%", thickness=0.8, color=colors.HexColor("#" + GOLD),
                                    spaceBefore=1, spaceAfter=7))
            index += 1
            continue
        if line.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            story.append(pdf_table(md_table_rows(table_lines), usable))
            story.append(Spacer(1, 6))
            continue
        heading_match = re.match(r"^(#{1,4})\s+(.*)$", line)
        if heading_match:
            level = min(len(heading_match.group(1)), 4)
            heading = heading_match.group(2).strip()
            styles = {1: h1, 2: h2, 3: h3, 4: h4}
            story.append(Paragraph(reportlab_inline(heading), styles[level]))
            index += 1
            continue
        if line.startswith("- "):
            story.append(Paragraph(reportlab_inline(line[2:]), bullet, bulletText="•"))
            index += 1
            continue
        if re.match(r"^\d+\.\s+", line):
            story.append(Paragraph(reportlab_inline(line), ordered))
            index += 1
            continue
        if line.startswith("> "):
            story.append(Paragraph(reportlab_inline(line[2:]), bullet))
            index += 1
            continue
        if "END-OF-SEMESTER EXAMINATION" in line or line.startswith("**Confidential:"):
            story.append(Paragraph(reportlab_inline(line), centered))
        else:
            style = guide_lead if guide and line.startswith("**(") else body
            story.append(Paragraph(reportlab_inline(line), style))
        index += 1

    def on_page(canvas, document):
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor("#" + GOLD))
        canvas.setLineWidth(0.65)
        canvas.line(left, page_height - 38, page_width - right, page_height - 38)
        canvas.setFont("ExamBody-Bold", 7.6)
        canvas.setFillColor(colors.HexColor("#" + MAROON))
        canvas.drawString(left, page_height - 29, "EXPLOITS UNIVERSITY  |  BICT 3202")
        label = "CONFIDENTIAL — EXAMINER COPY" if guide else f"PAPER {paper} — CANDIDATE COPY"
        canvas.drawRightString(page_width - right, page_height - 29, label)
        canvas.setStrokeColor(colors.HexColor("#D8C9C9"))
        canvas.line(left, 32, page_width - right, 32)
        canvas.setFont("ExamBody", 7.4)
        canvas.setFillColor(colors.HexColor("#" + GREY))
        canvas.drawString(left, 20, "Object-Oriented Analysis and Design · BICT 3202")
        canvas.drawRightString(page_width - right, 20, f"Page {document.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(str(destination), pagesize=A4, leftMargin=left, rightMargin=right,
                            topMargin=top, bottomMargin=bottom,
                            title=source.stem.replace("-", " "),
                            author="Exploits University — BICT 3202")
    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)


def main() -> None:
    for source in PAPER_SOURCES:
        if not source.exists():
            raise FileNotFoundError(source)
        folder = EXAM_OUTPUT / ("candidates" if is_candidate(source) else "examiner")
        folder.mkdir(parents=True, exist_ok=True)
        docx_path = folder / source.with_suffix(".docx").name
        pdf_path = folder / source.with_suffix(".pdf").name
        docx_build(source, docx_path)
        pdf_build(source, pdf_path)
        print(f"Built {docx_path.relative_to(ROOT)}")
        print(f"Built {pdf_path.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
