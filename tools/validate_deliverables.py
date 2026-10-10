"""Fail-closed regression checks for the reviewed Week 15/16 and exam outputs.

Run from any directory: python tools/validate_deliverables.py
Use --write-checksums only after a successful review of rebuilt deliverables.
"""
from __future__ import annotations

import argparse
import ast
import hashlib
import re
import sys
import zipfile
from pathlib import Path

import pymupdf
from docx import Document
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "output"
sys.path.insert(0, str(ROOT / "tools"))
from build_exam_docs import EXAM_OUTPUT, PAPER_SOURCES, inline_tokens, md_table_rows
from revision_assignments import ASSIGNMENTS, blocks


def require(condition, message):
    if not condition:
        raise AssertionError(message)


def normalise(text):
    return "".join(c.lower() for c in text if c.isalnum())


def word_text(path):
    doc = Document(path)
    return "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(
        c.text for table in doc.tables for row in table.rows for c in row.cells)


def expected_files():
    paths = []
    for week in (15, 16):
        for kind, ext in (("Module", "pdf"), ("Presentation", "pptx"),
                          ("Assignment", "pdf"), ("Assignment", "docx")):
            paths.append(OUTPUT / f"week{week}" / f"Week{week}_{kind}_BICT3202_OOAD.{ext}")
    for source in PAPER_SOURCES:
        folder = EXAM_OUTPUT / ("candidates" if "Paper-" in source.name else "examiner")
        paths.extend(folder / source.with_suffix(ext).name for ext in (".pdf", ".docx"))
    return sorted(paths)


def plain_markdown(text):
    return "".join(plain_markdown(value) if kind in ("bold", "italic") else value
                   for value, kind in inline_tokens(text))


def markdown_passages(source):
    for line in source.read_text(encoding="utf-8").splitlines():
        if not line.strip() or line.strip() == "---":
            continue
        if line.startswith("|"):
            for row in md_table_rows([line]):
                yield from (plain_markdown(c) for c in row if c)
        else:
            line = re.sub(r"^(?:#{1,4}\s+|-\s+|\d+\.\s+|>\s+)", "", line)
            yield plain_markdown(line)


def check_passages(passages, text, label):
    haystack = normalise(text)
    for passage in passages:
        require(normalise(passage) in haystack, f"{label}: missing content: {passage}")


def check_exam_marks(text):
    sections = re.split(r"^### Question (\d+) .*?\((\d+) marks\)\s*$", text, flags=re.M)
    require(len(sections) == 16, "Exam must contain exactly Questions 1–5")
    totals = []
    for i in range(1, len(sections), 3):
        q, advertised, body = sections[i:i+3]
        expected = 40 if q == "1" else 20
        marks = [int(m) for m in re.findall(r"\*\*\[(\d+) marks\]\*\*", body)]
        labels = re.findall(r"\*\*\(([a-z])\)\*\*", body)
        require(labels == [chr(97+j) for j in range(len(labels))], f"Q{q}: bad subpart order")
        require(len(marks) == len(labels) and sum(marks) == int(advertised) == expected,
                f"Q{q}: marks do not match heading/subparts")
        totals.append(expected)
    require(totals == [40, 20, 20, 20, 20], "Invalid exam question sequence")
    require(totals[0] + sum(totals[1:4]) == 100, "Candidate selection must total 100")
    require("any THREE" in text and "Question 1 (compulsory)" in text, "Missing selection rule")


def check_assignment_marks(data):
    require(sum(m for _, m in data["short"]) == 20, "Assignment Section A must total 20")
    require([sum(m for _, m in questions) for _, questions in data["application"]] == [12, 8],
            "Assignment Section B must total 12 + 8")
    require(sum(m for _, m in data["case"][1]) == 10, "Assignment Section C must total 10")


def check_pdf(path):
    with pymupdf.open(path) as doc:
        require(doc.page_count >= 2, f"{path.name}: unexpectedly short PDF")
        require(not doc.is_encrypted, f"{path.name}: encrypted PDF")
        for page in doc:
            require(len(page.get_text().strip()) > 70, f"{path.name}: blank page {page.number+1}")
            require(abs(page.rect.width - 595.276) < 1 and abs(page.rect.height - 841.89) < 1,
                    f"{path.name}: expected A4 page")
            require("BICT 3202" in page.get_text(), f"{path.name}: missing course/page label")
            require("\ufffd" not in page.get_text(), f"{path.name}: replacement glyph")
            for block in page.get_text("dict")["blocks"]:
                for line in block.get("lines", []):
                    for span in line["spans"]:
                        x0, y0, x1, y1 = span["bbox"]
                        require(x0 >= -0.5 and y0 >= -0.5 and
                                x1 <= page.rect.width + 0.5 and y1 <= page.rect.height + 0.5,
                                f"{path.name} p{page.number+1}: text clipped: {span['text']}")
        return "\n".join(p.get_text() for p in doc), doc.page_count


def check_office(path):
    with zipfile.ZipFile(path) as archive:
        require(archive.testzip() is None, f"{path.name}: corrupt archive")
        # Every image embedded in Office must be decodable, not just named in XML.
        for name in archive.namelist():
            if "/media/" in name:
                import io
                with Image.open(io.BytesIO(archive.read(name))) as image:
                    image.verify()
    if path.suffix == ".docx":
        doc = Document(path)
        width = doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin
        for table in doc.tables:
            require(sum(c.width for c in table.columns) <= width + 1000,
                    f"{path.name}: table exceeds text width")
        return word_text(path)
    deck = Presentation(path)
    require(len(deck.slides) == 26, f"{path.name}: expected 26 slides")
    for i, slide in enumerate(deck.slides):
        for shape in slide.shapes:
            require(shape.left >= 0 and shape.top >= 0 and
                    shape.left + shape.width <= deck.slide_width + 1000 and
                    shape.top + shape.height <= deck.slide_height + 1000,
                    f"{path.name} slide {i+1}: off-slide shape {shape.name}")
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and i > 0:
                require(shape.top >= 1.26*914400 and shape.top+shape.height <= 6.45*914400+1000,
                        f"{path.name} slide {i+1}: picture overlaps header/takeaway/footer")
    return "\n".join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)


def check_links(writing_manifest=False):
    for source in (ROOT / "README.md", ROOT / "exams/README.md", OUTPUT / "README.md"):
        require(source.exists(), f"Missing {source}")
        for link in re.findall(r"\]\(([^)]+)\)", source.read_text()):
            if "://" in link or link.startswith("#"):
                continue
            target = source.parent / link.split("#")[0]
            if writing_manifest and target.resolve() == (OUTPUT / "checksums.sha256").resolve():
                continue  # created only after the other acceptance checks pass
            require(target.exists(), f"{source}: broken link {link}")


def syntax_check():
    for path in (ROOT / "tools").glob("*.py"):
        ast.parse(path.read_text(encoding="utf-8"), filename=str(path))


def manifest_text():
    return "".join(f"{hashlib.sha256(p.read_bytes()).hexdigest()}  {p.relative_to(ROOT).as_posix()}\n"
                   for p in expected_files())


def validate(write_checksums=False):
    syntax_check()
    paths = expected_files()
    actual = sorted(p for folder in (OUTPUT / "week15", OUTPUT / "week16", EXAM_OUTPUT)
                    for p in folder.rglob("*") if p.suffix in (".pdf", ".docx", ".pptx"))
    require(paths == actual, "Missing deliverables or unexpected duplicate distribution files")
    require(not list((ROOT / "exams").glob("*.pdf")) and not list((ROOT / "exams").glob("*.docx")),
            "Generated exams must live only in output/exams")
    texts, counts = {}, {}
    for path in paths:
        require(path.exists(), f"Missing {path}")
        if path.suffix == ".pdf":
            texts[path], counts[path] = check_pdf(path)
        else:
            texts[path] = check_office(path)
    for source in PAPER_SOURCES:
        candidate = "Paper-" in source.name
        folder = EXAM_OUTPUT / ("candidates" if candidate else "examiner")
        if candidate:
            check_exam_marks(source.read_text())
        for ext in (".pdf", ".docx"):
            path = folder / source.with_suffix(ext).name
            text = texts[path]
            check_passages(markdown_passages(source), text, path.name)
            require("`" not in text and "**" not in text, f"{path.name}: raw Markdown leaked")
            if ext == ".pdf":
                require(("CONFIDENTIAL — EXAMINER COPY" in text) != candidate,
                        f"{path.name}: wrong PDF distribution label")
                if candidate:
                    require("CANDIDATE COPY" in text, f"{path.name}: missing candidate label")
            else:
                header = Document(path).sections[0].header.paragraphs[0].text
                require(("CONFIDENTIAL" in header) != candidate, f"{path.name}: wrong Word distribution label")
            if candidate:
                require("proposed" in text and "Faculty approval" in text, "Duration must be flagged as provisional")
        print(f"PASS {source.stem}: complete PDF/DOCX content, labels and marks")
    for week in (15, 16):
        data = ASSIGNMENTS[week]
        check_assignment_marks(data)
        passages = []
        for kind, value in blocks(week):
            passages.extend((cell for row in value for cell in row if cell) if kind == "table" else [value])
        for ext in ("pdf", "docx"):
            path = OUTPUT / f"week{week}" / f"Week{week}_Assignment_BICT3202_OOAD.{ext}"
            check_passages(passages, texts[path], path.name)
        module = OUTPUT / f"week{week}" / f"Week{week}_Module_BICT3202_OOAD.pdf"
        require(counts[module] >= 15 and "Contents" in texts[module] and "References" in texts[module],
                f"Week {week}: missing module sections")
        deck = OUTPUT / f"week{week}" / f"Week{week}_Presentation_BICT3202_OOAD.pptx"
        key_terms = ("reserveSeat", "transaction", "confirmation") if week == 15 else ("revision", "examination")
        for text in (texts[module], texts[deck]):
            check_passages(key_terms, text, f"Week {week} key themes")
        print(f"PASS Week {week}: module {counts[module]} pages, 26 slides, aligned 50-mark PDF/DOCX assignment")
    check_links(write_checksums)
    manifest = OUTPUT / "checksums.sha256"
    if write_checksums:
        manifest.write_text(manifest_text())
    else:
        require(manifest.exists() and manifest.read_text() == manifest_text(), "Checksum manifest mismatch")
    print(f"PASS all {len(paths)} documents: structure, source coverage, marks, page/slide bounds, links and checksums")
    return paths, counts


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--write-checksums", action="store_true")
    args = parser.parse_args()
    try:
        validate(args.write_checksums)
    except (AssertionError, OSError, ValueError) as error:
        print(f"FAIL: {error}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
